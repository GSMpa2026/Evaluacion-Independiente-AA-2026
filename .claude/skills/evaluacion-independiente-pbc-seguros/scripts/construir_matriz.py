#!/usr/bin/env python3
"""Construye la matriz de riesgo GSM sobre la plantilla Excel original.
Uso: python3 construir_matriz.py plantilla.xlsx salida.xlsx   (lee matrix_data.json del generador)
Conserva hojas Variables/Riesgos/Matriz de Riesgos, fórmulas y formatos condicionales; amplía rangos a la fila 120,
agrega columnas K-O y resumen por componente. Después: recalcular con LibreOffice (recalc.py de la skill xlsx).
Ajusta los textos de fecha de las filas 1-2 (ws[A2], var[A2]) para cada versión."""
import json, copy, re
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

import sys
SRC = sys.argv[1] if len(sys.argv) > 1 else 'plantilla_matriz_riesgo_GSM.xlsx'
OUT = sys.argv[2] if len(sys.argv) > 2 else 'Matriz_Riesgo_GSM.xlsx'
D = json.load(open('matrix_data.json'))
wb = openpyxl.load_workbook(SRC)
ws = wb['Riesgos']; mx = wb['Matriz de Riesgos']; var = wb['Variables']
LAST = 120

# --- encabezados de fecha ---
ws['A1'].value = 'ASEGURADORA ANCÓN, S.A. – MATRIZ DE RIESGO GSM – EVALUACIÓN INDEPENDIENTE 2024-2025'
ws['A2'].value = 'Período: 1 julio 2024 – 30 junio 2025  |  Elaborado: José A. Serrano   |  Versión inicial: 02 de abril de 2026  |  Actualizada: octubre de 2026 (corte de evidencia: cuadro de control octubre 2026)'
var['A2'].value = 'Evaluación Independiente 2024-2025  |  Preparado por: José A. Serrano   |  02 abril 2026 (actualizado 02 octubre 2026)'
ws.unmerge_cells('A1:J1'); ws.unmerge_cells('A2:J2')
ws.merge_cells('A1:O1'); ws.merge_cells('A2:O2')
for c in 'KLMNO':
    for r in (1, 2):
        ws[f'{c}{r}']._style = copy.copy(ws[f'J{r}']._style)

# --- estilos base ---
hdr = ws['J3']; st_w = {c: copy.copy(ws[f'{c}4']._style) for c in 'ABCDEFGHIJ'}; st_g = {c: copy.copy(ws[f'{c}5']._style) for c in 'ABCDEFGHIJ'}
for c in 'KLMNO':
    st_w[c] = copy.copy(ws['C4']._style); st_g[c] = copy.copy(ws['C5']._style)
st_w['M'] = copy.copy(ws['F4']._style); st_g['M'] = copy.copy(ws['F5']._style)
ws['D3'].value = 'OBSERVACIÓN / NOTA GSM\n(actualizada oct-2026)'
for c, t in zip('KLMNO', ['COMPONENTE', 'BASE NORMATIVA / REFERENCIA', 'PUNTAJE\nP×I', 'PLAN DE ACCIÓN / FASE', 'ACCIÓN REQUERIDA SEGÚN NIVEL']):
    ws[f'{c}3'].value = t; ws[f'{c}3']._style = copy.copy(hdr._style)
for c, w in zip('ABCDEFGHIJKLMNO', [62, 18, 30, 40, 14, 8, 14, 8, 10, 16, 20, 34, 9, 16, 44]):
    ws.column_dimensions[c].width = w

# --- limpiar filas previas ---
for r in range(4, LAST + 1):
    for c in range(1, 16):
        ws.cell(r, c).value = None

# --- filas ---
PF = {'Fase 1': ('Ocasional', 'Moderado'), 'Fase 2': ('Posible', 'Moderado'), 'Fase 3': ('Posible', 'Menor')}
rows = []
for h in D['hallazgos']:
    rows.append(dict(A=f"{h['id']} · {h['short']}. {h['text']}", B='Brecha', C=h['docs'], D=h['nota'], E=h['p'], G=h['i'], K=h['comp'], L=h['base'], N=h['pa']))
for o in D['omx']:
    rows.append(dict(A=f"{o['id']} · {o['tema']}. {o['desc']}", B='Oportunidad de Mejora', C=o['docs'], D=o['accion'], E='Posible', G='Menor', K='Período 2024-2025', L='Acuerdo 03-2022 (buena práctica)', N='Período'))
for m in D['rm']:
    p, i = PF[m['fase']]
    rows.append(dict(A=f"{m['id']} · {m['tema']}. {m['mejora']}", B='Oportunidad de Mejora', C=m['comp'], D=f"Qué cambia: {m['cambia']} Situación en Ancón: {m['sit']}", E=p, G=i, K='Hoja de ruta Acuerdo 07-2025', L='Acuerdo SSRP 07-2025 (fuera del alcance; valor agregado)', N=m['fase']))
assert len(rows) <= LAST - 3
for k, row in enumerate(rows):
    r = 4 + k; st = st_w if k % 2 == 0 else st_g
    for c in 'ABCDEFGHIJKLMNO':
        ws[f'{c}{r}']._style = copy.copy(st[c])
    for c, v in row.items(): ws[f'{c}{r}'].value = v
    ws[f'F{r}'] = f'=IFERROR(VLOOKUP(E{r},Variables!$A$6:$E$10,5,0),"")'
    ws[f'H{r}'] = f'=IFERROR(VLOOKUP(G{r},Variables!$A$14:$E$18,5,0),"")'
    ws[f'I{r}'] = f'=IF(B{r}="Oportunidad de Mejora","OM",IFERROR(CONCATENATE(F{r},":",H{r}),""))'
    ws[f'J{r}'] = f'=IFERROR(VLOOKUP(I{r},Variables!$A$22:$B$47,2,FALSE()),"Por definir")'
    ws[f'M{r}'] = f'=IF(I{r}="OM","N/A",IFERROR(F{r}*H{r},""))'
    ws[f'O{r}'] = f'=IFERROR(VLOOKUP(J{r},Variables!$A$51:$B$55,2,FALSE()),"")'
    for c in 'ACDKLNO':
        ws[f'{c}{r}'].alignment = Alignment(wrap_text=True, vertical='top')
    for c in 'BEFGHIJM':
        ws[f'{c}{r}'].alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    ws.row_dimensions[r].height = max(60, min(170, 13 * (len(row['A']) // 70 + 2)))
last_row = 3 + len(rows)

# --- formato condicional y validaciones ---
old = []
for cf in ws.conditional_formatting:
    old.append((str(cf.sqref), list(cf.rules)))
ws.conditional_formatting = type(ws.conditional_formatting)()
for sq, rules in old:
    new = re.sub(r'(\D)50\b', lambda m: m.group(1) + str(LAST), sq).replace('J50', f'J{LAST}').replace('H50', f'H{LAST}')
    if sq.startswith('A4:J'): new = f'A4:O{LAST}'
    for rl in rules: ws.conditional_formatting.add(new, rl)
ws.data_validations.dataValidation = []
dv = DataValidation(type='list', formula1='"Brecha,Oportunidad de Mejora"', allow_blank=True); dv.add(f'B4:B{LAST}'); ws.add_data_validation(dv)
dvp = DataValidation(type='list', formula1='=Variables!$A$6:$A$10', allow_blank=True); dvp.add(f'E4:E{LAST}'); ws.add_data_validation(dvp)
dvi = DataValidation(type='list', formula1='=Variables!$A$14:$A$18', allow_blank=True); dvi.add(f'G4:G{LAST}'); ws.add_data_validation(dvi)
ws.auto_filter.ref = f'A3:O{last_row}'
ws.freeze_panes = 'B4'

# --- hoja Matriz de Riesgos: ampliar rangos ---
for row in mx.iter_rows():
    for c in row:
        if isinstance(c.value, str) and c.value.startswith('='):
            c.value = c.value.replace('$50)', f'${LAST})').replace('$50,', f'${LAST},')
mx['A2'].value = f'Los valores de cada celda son generados automáticamente desde la hoja Riesgos  |  COUNTIF(Riesgos!$I$4:$I${LAST})  |  Evaluación Independiente 2024-2025 (actualizada octubre 2026)'

# resumen por componente
thin = Side(style='thin', color='BFBFBF'); bd = Border(left=thin, right=thin, top=thin, bottom=thin)
r0 = 27
mx.merge_cells(start_row=r0, start_column=1, end_row=r0, end_column=8)
mx.cell(r0, 1).value = 'RESUMEN POR COMPONENTE – BRECHAS DEL PERÍODO (julio 2024 – junio 2025)'
mx.cell(r0, 1)._style = copy.copy(mx['A18']._style)
heads = ['Componente', '', '', 'CRÍTICO', 'MEDIO-ALTO', 'MEDIO', 'BAJO', 'TOTAL']
hfill = {'CRÍTICO': 'FFC00000', 'MEDIO-ALTO': 'FFED7D31', 'MEDIO': 'FFFFD966', 'BAJO': 'FF70AD47'}
for j, t in enumerate(heads, start=1):
    c = mx.cell(r0 + 1, j); c.value = t if t else None
    c.font = Font(name='Arial', bold=True, size=9, color='FFFFFF' if t in ('CRÍTICO', 'MEDIO-ALTO', 'Componente', 'TOTAL') else '000000')
    c.fill = PatternFill('solid', fgColor=hfill.get(t, 'FF1F3864')); c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); c.border = bd
mx.merge_cells(start_row=r0 + 1, start_column=1, end_row=r0 + 1, end_column=3)
comps = []
for h in D['hallazgos']:
    if h['comp'] not in comps: comps.append(h['comp'])
for k, comp in enumerate(comps):
    r = r0 + 2 + k
    mx.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    mx.cell(r, 1).value = comp
    mx.cell(r, 1).font = Font(name='Arial', size=9, bold=True); mx.cell(r, 1).border = bd
    for j, lev in zip(range(4, 8), ['CRÍTICO', 'MEDIO-ALTO', 'MEDIO', 'BAJO']):
        c = mx.cell(r, j); c.value = f'=COUNTIFS(Riesgos!$K$4:$K${LAST},$A{r},Riesgos!$J$4:$J${LAST},"{lev}")'
        c.font = Font(name='Arial', size=9); c.alignment = Alignment(horizontal='center'); c.border = bd
    c = mx.cell(r, 8); c.value = f'=SUM(D{r}:G{r})'; c.font = Font(name='Arial', size=9, bold=True); c.alignment = Alignment(horizontal='center'); c.border = bd
rt = r0 + 2 + len(comps)
mx.merge_cells(start_row=rt, start_column=1, end_row=rt, end_column=3)
mx.cell(rt, 1).value = 'TOTAL'; mx.cell(rt, 1).font = Font(name='Arial', size=9, bold=True); mx.cell(rt, 1).fill = PatternFill('solid', fgColor='FFD6DCE4')
for j in range(4, 9):
    col = get_column_letter(j); c = mx.cell(rt, j); c.value = f'=SUM({col}{r0+2}:{col}{rt-1})'
    c.font = Font(name='Arial', size=9, bold=True); c.alignment = Alignment(horizontal='center'); c.fill = PatternFill('solid', fgColor='FFD6DCE4'); c.border = bd
mx.cell(rt + 2, 1).value = 'Nota: las oportunidades de mejora (OM) no se califican como brechas. Las RM corresponden a la hoja de ruta hacia el Acuerdo SSRP 07-2025, que está fuera del alcance de la evaluación del período y se incluye como valor agregado.'
mx.cell(rt + 2, 1).font = Font(name='Arial', size=8, italic=True, color='595959')
mx.merge_cells(start_row=rt + 2, start_column=1, end_row=rt + 2, end_column=8)
mx.cell(rt + 2, 1).alignment = Alignment(wrap_text=True); mx.row_dimensions[rt + 2].height = 30

wb.save(OUT)
print('rows', len(rows), 'last', last_row, 'comps', comps)
